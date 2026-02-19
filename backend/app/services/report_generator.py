"""
Word (.docx) report generator using python-docx.
Generates the report into an in-memory BytesIO buffer — nothing is saved to disk.
"""

import io
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH

from app.models.models import Application, Inspection

# Human-readable labels for metric fields
METRIC_LABELS = {
    "heating_consumption": ("Потребление тепла", "Гкал"),
    "electricity_consumption": ("Потребление электроэнергии", "кВт·ч"),
    "water_consumption": ("Потребление воды", "м³"),
    "gas_consumption": ("Потребление газа", "м³"),
    "wall_thickness_mm": ("Толщина стен", "мм"),
    "window_type": ("Тип окон", ""),
    "insulation_type": ("Тип утепления", ""),
    "thermal_resistance": ("Термическое сопротивление", "м²·°C/Вт"),
    "air_tightness": ("Воздухопроницаемость", "м³/(ч·м²)"),
    "indoor_temperature": ("Температура внутри", "°C"),
    "outdoor_temperature": ("Температура снаружи", "°C"),
}


def generate_report(application: Application) -> io.BytesIO:
    doc = Document()

    # ── Styles ─────────────────────────────────────────────────────────
    style = doc.styles["Normal"]
    font = style.font
    font.name = "Times New Roman"
    font.size = Pt(12)

    # ── Title ──────────────────────────────────────────────────────────
    title = doc.add_heading("Отчёт по энергетическому обследованию", level=1)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.add_paragraph(f"Заявка №{application.id}")
    doc.add_paragraph(
        f"Дата создания заявки: {application.created_at.strftime('%d.%m.%Y')}"
    )
    doc.add_paragraph(f"Тип услуги: {application.service_type}")

    # ── Object info ────────────────────────────────────────────────────
    obj = application.audit_object
    if obj:
        doc.add_heading("1. Объект обследования", level=2)
        doc.add_paragraph(f"Адрес: {obj.address}")
        doc.add_paragraph(f"Тип объекта: {obj.object_type}")
        if obj.total_area:
            doc.add_paragraph(f"Общая площадь: {obj.total_area} м²")
        if obj.floors:
            doc.add_paragraph(f"Этажность: {obj.floors}")
        if obj.year_built:
            doc.add_paragraph(f"Год постройки: {obj.year_built}")
        if obj.description:
            doc.add_paragraph(f"Описание: {obj.description}")

    # ── Inspections ────────────────────────────────────────────────────
    inspections: list[Inspection] = application.inspections
    if inspections:
        doc.add_heading("2. Результаты обследований", level=2)

        for idx, insp in enumerate(inspections, start=1):
            engineer_name = insp.engineer.full_name if insp.engineer else f"ID {insp.engineer_id}"
            doc.add_heading(
                f"2.{idx}. Обследование — {engineer_name}", level=3
            )
            doc.add_paragraph(
                f"Статус: {'Отправлено' if insp.status.value == 'submitted' else 'Черновик'}"
            )
            if insp.submitted_at:
                doc.add_paragraph(
                    f"Дата отправки: {insp.submitted_at.strftime('%d.%m.%Y %H:%M')}"
                )

            # Metrics table
            table = doc.add_table(rows=1, cols=3)
            table.style = "Table Grid"
            hdr = table.rows[0].cells
            hdr[0].text = "Метрика"
            hdr[1].text = "Значение"
            hdr[2].text = "Ед. изм."

            for field, (label, unit) in METRIC_LABELS.items():
                value = getattr(insp, field, None)
                if value is not None:
                    row = table.add_row().cells
                    row[0].text = label
                    row[1].text = str(value)
                    row[2].text = unit

            # Extra metrics from JSONB
            if insp.extra_metrics:
                for key, value in insp.extra_metrics.items():
                    row = table.add_row().cells
                    row[0].text = key
                    row[1].text = str(value)
                    row[2].text = ""

            if insp.notes:
                doc.add_paragraph(f"Примечания: {insp.notes}")

            doc.add_paragraph("")  # spacing

    # ── Summary section (placeholder for manual completion) ────────────
    doc.add_heading("3. Выводы и рекомендации", level=2)
    doc.add_paragraph(
        "[Данный раздел заполняется инженером вручную после генерации отчёта]"
    )

    # ── Write to buffer ────────────────────────────────────────────────
    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf
